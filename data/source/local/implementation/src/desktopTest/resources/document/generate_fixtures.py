# This file is part of Campfire.
# Copyright (c) Pandula Péter 2017-2026.
# https://github.com/pandulapeter/campfire
#
# This Source Code Form is subject to the terms of the Mozilla Public License, v. 2.0.
# If a copy of the MPL was not distributed with this file, You can obtain one at
# https://mozilla.org/MPL/2.0/.

"""Independent generator fixtures, using the public-domain song Amazing Grace.

Run once with python-docx, ReportLab and pypdf; the generated files are committed,
not rebuilt by tests. LibreOffice and Chrome exports are separate manual commands
documented in README.md so their actual producers remain visible.
"""

from pathlib import Path
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parent
LYRIC = "Amazing grace how sweet the sound"
CHORDS = "G             C             G"


def word():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for name in ("Normal", "Title", "Heading 1"):
        font = doc.styles[name].font
        font.name = "Courier New"
        font.size = Pt(12 if name == "Normal" else 22 if name == "Title" else 12)
        font.color.rgb = RGBColor(0, 0, 0)
        for border in list(doc.styles[name].element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
        for fonts in doc.styles[name].element.iter(qn("w:rFonts")):
            for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                fonts.attrib.pop(qn("w:" + attribute), None)
        for color in doc.styles[name].element.iter(qn("w:color")):
            color.attrib.pop(qn("w:themeColor"), None)
    doc.add_paragraph("Amazing Grace", "Title")
    doc.add_paragraph("by John Newton")
    doc.add_paragraph("Verse 1", "Heading 1")
    doc.add_paragraph(CHORDS)
    doc.add_paragraph(LYRIC)
    path = ROOT / "python-docx.docx"
    doc.save(path)
    with ZipFile(path) as original:
        entries = [(name, original.read(name)) for name in original.namelist()]
    with ZipFile(path, "w", ZIP_DEFLATED) as stable:
        for name, data in entries:
            info = ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            stable.writestr(info, data)


def sheet(pdf, x=50, y=700, section="Verse 1", font="Courier", title=True):
    if title:
        pdf.setFont("Helvetica-Bold", 22)
        pdf.drawString(x, y + 50, "Amazing Grace")
        pdf.setFont("Helvetica", 10)
        pdf.drawString(x, y + 30, "by John Newton")
    pdf.setFont(font, 12)
    pdf.drawString(x, y, section)
    pdf.drawString(x, y - 20, CHORDS)
    pdf.drawString(x, y - 34, LYRIC)


def pdfs():
    pdf = canvas.Canvas(str(ROOT / "reportlab.pdf"), pagesize=(612, 792), pageCompression=1, invariant=1)
    sheet(pdf)
    pdf.save()
    pdfmetrics.registerFont(TTFont("Unicode", "/System/Library/Fonts/Supplemental/Arial.ttf"))
    pdf = canvas.Canvas(str(ROOT / "reportlab-unicode.pdf"), pagesize=(612, 792), pageCompression=1, invariant=1)
    pdf.setFont("Unicode", 12)
    pdf.drawString(50, 700, "\u00c1rv\u00edzt\u0171r\u0151 t\u00fck\u00f6rf\u00far\u00f3g\u00e9p \u03a9 \u0416")
    pdf.save()
    pdf = canvas.Canvas(str(ROOT / "two-column.pdf"), pagesize=(800, 792), pageCompression=1, invariant=1)
    sheet(pdf, x=40)
    sheet(pdf, x=420, section="Verse 2", title=False)
    pdf.save()
    pdf = canvas.Canvas(str(ROOT / "songbook.pdf"), pagesize=(612, 792), pageCompression=1, invariant=1)
    sheet(pdf)
    pdf.showPage()
    pdf.setFont("Helvetica-Bold", 22)
    pdf.drawString(50, 750, "Amazing Grace Second Sheet")
    sheet(pdf, section="Verse 2", title=False)
    pdf.save()
    reader = PdfReader(ROOT / "reportlab.pdf")
    writer = PdfWriter()
    writer.append(reader)
    writer.encrypt("")
    writer.write(ROOT / "protected.pdf")


if __name__ == "__main__":
    word()
    pdfs()
